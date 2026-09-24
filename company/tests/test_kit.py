# -*- coding: utf-8 -*-
"""
公司工具包回归测试

改了 company 文件夹、升级终端镜像或 Open WebUI 之后跑一次，全部通过再同步 NAS：
    本机：本机测试工具.bat → 菜单 C
    NAS ：Container Manager → 容器 ai-term-zc → 终端机 → python3 /opt/company/tests/test_kit.py

只在 /tmp 下建临时文件，不碰部门文件夹。缺 LibreOffice 等组件的项目会标「跳过」。
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
COMPANY = os.path.dirname(HERE)
KIT = os.path.join(COMPANY, "kit")
sys.path.insert(0, KIT)
PY = sys.executable
MARK = "内部文件，禁止外传"


def run(*args, cwd=None):
    return subprocess.run([PY, *args], capture_output=True, text=True, cwd=cwd, timeout=600)


def make_dept(tmp, name="测试部"):
    root = os.path.join(tmp, name)
    for s in ("01_原始数据", "02_输出", "03_周报", "04_技能", "05_模板"):
        os.makedirs(os.path.join(root, s))
    return root


def sample_xlsx(path, rows=None, cols=("仓库", "金额")):
    import pandas as pd
    rows = rows or [("A仓", 10), ("B仓", 20), ("A仓", 30), ("C仓", 5.5)]
    pd.DataFrame(rows, columns=list(cols)).to_excel(path, index=False)
    return path


SCRIPT = '''import sys; sys.path.insert(0, %r)
from runio import io_args, save_summary
import pandas as pd
inputs, out = io_args([%r], %r)
df = pd.read_excel(inputs[0])
save_summary({"总金额": df["金额"].sum(), "行数": len(df), "按仓库": df.groupby("仓库")["金额"].sum().to_dict()})
print("完成", df["金额"].sum())
'''


class T01_Runio(unittest.TestCase):
    def test_io_and_summary(self):
        tmp = tempfile.mkdtemp()
        try:
            f = sample_xlsx(os.path.join(tmp, "in.xlsx"))
            out = os.path.join(tmp, "out")
            script = os.path.join(tmp, "script.py")
            open(script, "w", encoding="utf-8").write(SCRIPT % (KIT, f, out))
            r = run(script)
            self.assertEqual(r.returncode, 0, r.stderr)
            doc = json.load(open(os.path.join(out, "结果摘要.json"), encoding="utf-8"))
            self.assertAlmostEqual(doc["总金额"], 65.5)
            self.assertEqual(doc["行数"], 4)
            self.assertEqual(doc["按仓库"]["A仓"], 40)
            self.assertEqual(doc["_输入结构"]["in.xlsx"]["Sheet1"], ["仓库", "金额"])
            # 换输入重跑
            f2 = sample_xlsx(os.path.join(tmp, "in2.xlsx"), rows=[("A仓", 1)])
            out2 = os.path.join(tmp, "out2")
            r = run(script, "--input", f2, "--out", out2)
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertEqual(json.load(open(os.path.join(out2, "结果摘要.json"), encoding="utf-8"))["总金额"], 1)
        finally:
            shutil.rmtree(tmp)


class T02_Dept(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.root = make_dept(self.tmp)
        self.dept = os.path.join(KIT, "dept.py")

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def d(self, *args):
        return run(self.dept, "--dept-dir", self.root, *args)

    def test_notes(self):
        r = self.d("note", "--section", "字段口径", "--by", "张三", "金额 = 含税金额")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.d("note", "--section", "术语", "--by", "张三", "A仓 = 主仓库")
        r = self.d("note", "--section", "字段口径", "--by", "李四", "金额 = 含税金额")
        self.assertIn("未重复写入", r.stdout)
        text = open(os.path.join(self.root, "05_模板", "部门说明.md"), encoding="utf-8").read()
        self.assertIn(MARK, text)
        self.assertEqual(text.count("金额 = 含税金额"), 1)
        self.assertLess(text.index("A仓 = 主仓库"), text.index("金额 = 含税金额"))  # 按小节放
        r = self.d("start")
        self.assertIn("2 条", r.stdout)
        self.assertIn("A仓 = 主仓库", r.stdout)
        self.assertNotEqual(self.d("note", "--section", "不存在", "x").returncode, 0)

    def test_notes_conflict_replace_cap(self):
        self.d("note", "--section", "字段口径", "--by", "张三", "入库金额 = 数量 × 含税单价")
        r = self.d("note", "--section", "字段口径", "--by", "李四", "入库金额 = 数量 × 不含税单价")
        self.assertEqual(r.returncode, 3)                      # 同名口径：不写入，列出旧条目
        self.assertIn("已经有口径", r.stdout)
        r = self.d("note", "--replace", "入库金额", "--by", "李四", "入库金额 = 数量 × 不含税单价")
        self.assertEqual(r.returncode, 0, r.stderr)
        text = open(os.path.join(self.root, "05_模板", "部门说明.md"), encoding="utf-8").read()
        self.assertIn("不含税单价", text)
        self.assertNotIn("× 含税单价", text)
        log = open(os.path.join(self.root, "05_模板", "部门说明_修改记录.md"), encoding="utf-8").read()
        self.assertIn("× 含税单价", log)                       # 旧内容留痕
        r = run(self.dept, "--dept-dir", self.root, "note", "--section", "其他", "x" * 50)
        env = dict(os.environ, AI_NOTES_MAX_CHARS="100")
        r = subprocess.run([PY, self.dept, "--dept-dir", self.root, "note", "--section", "其他", "y" * 60],
                           capture_output=True, text=True, env=env)
        self.assertEqual(r.returncode, 4)                      # 超过篇幅上限：不写入
        self.assertIn("篇幅上限", r.stdout)

    def test_notes_list_del_review(self):
        for i in range(6):
            self.d("note", "--section", "术语", "--by", "张三", f"词{i} = 含义{i}")
        self.assertIn("没复核", self.d("start").stdout)          # 5 条以上、从未复核 → 提醒
        r = self.d("note-list")
        self.assertIn("#6", r.stdout)
        r = self.d("note-del", "2", "3", "--by", "张三")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("4 条", r.stdout)
        self.d("note-review", "--by", "张三")
        out = self.d("start").stdout
        self.assertIn("最近复核 20", out)
        self.assertNotIn("没复核", out)

    def test_recipe_cycle(self):
        src_in = sample_xlsx(os.path.join(self.root, "01_原始数据", "入库_0923.xlsx"))
        task = os.path.join(self.root, "02_输出", "20260924_张三_入库汇总")
        os.makedirs(task)
        open(os.path.join(task, "script.py"), "w", encoding="utf-8").write(SCRIPT % (KIT, src_in, task))
        self.assertEqual(run(os.path.join(task, "script.py")).returncode, 0)
        r = self.d("recipe-save", "--name", "入库汇总", "--from", task, "--desc", "按仓库汇总入库金额", "--by", "张三")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        rd = os.path.join(self.root, "04_技能", "入库汇总")
        for f in ("SKILL.md", "script.py", "验收.json", os.path.join("样例", "入库_0923.xlsx")):
            self.assertTrue(os.path.isfile(os.path.join(rd, f)), f)
        self.assertIn("入库汇总", self.d("start").stdout)
        # 新文件格式一致 → 能跑
        new = sample_xlsx(os.path.join(self.root, "01_原始数据", "入库_0930.xlsx"), rows=[("A仓", 1), ("B仓", 2)])
        r = self.d("recipe-run", "入库汇总", "--input", new, "--user", "李四")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn('"总金额": 3', r.stdout)
        # 列名变了 → 拒绝运行
        bad = sample_xlsx(os.path.join(self.root, "01_原始数据", "bad.xlsx"), cols=("库房", "金额"))
        r = self.d("recipe-run", "入库汇总", "--input", bad)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("格式不一致", r.stdout + r.stderr)
        # 回归：通过；篡改验收值 → 不一致
        self.assertEqual(self.d("recipe-check").returncode, 0)
        acc = os.path.join(rd, "验收.json")
        doc = json.load(open(acc, encoding="utf-8"))
        doc["总金额"] = 999
        json.dump(doc, open(acc, "w", encoding="utf-8"), ensure_ascii=False)
        r = self.d("recipe-check")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("不一致", r.stdout)
        # 覆盖更新：旧版进 历史/
        r = self.d("recipe-save", "--name", "入库汇总", "--from", task, "--desc", "v2", "--overwrite")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(len(os.listdir(os.path.join(rd, "历史"))), 1)
        self.assertEqual(self.d("recipe-check").returncode, 0)

    def test_recipe_requires_runio(self):
        task = os.path.join(self.root, "02_输出", "x")
        os.makedirs(task)
        open(os.path.join(task, "script.py"), "w", encoding="utf-8").write("print(1)\n")
        r = self.d("recipe-save", "--name", "x", "--from", task, "--desc", "x")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("可重跑", r.stdout + r.stderr)
        self.assertFalse(os.path.exists(os.path.join(self.root, "04_技能", "x")))


class T03_Peek(unittest.TestCase):
    def test_peek(self):
        import pandas as pd
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, "表头在第3行.xlsx")
            df = pd.DataFrame([["入库明细", None], [None, None], ["物料", "数量"], ["M1", 3], ["M2", "4"]])
            df.to_excel(p, index=False, header=False)
            r = run(os.path.join(KIT, "peek.py"), p)
            self.assertEqual(r.returncode, 0, r.stderr)
            self.assertLessEqual(len(r.stdout.splitlines()), 80)
            self.assertIn("3", r.stdout)
        finally:
            shutil.rmtree(tmp)


class T04_ChartsTables(unittest.TestCase):
    def test_charts(self):
        import charts
        tmp = tempfile.mkdtemp()
        try:
            outs = [
                charts.line(["W1", "W2", "W3"], {"金额": [1, 2, 3]}, "金额 · 测试", "万元", os.path.join(tmp, "a.png")),
                charts.bar([f"类{i}" for i in range(12)], list(range(12)), "多类别 · 测试", "个", os.path.join(tmp, "b.png")),
                charts.bar_compare(["A", "B"], {"上周": [1, 2], "本周": [2, 1]}, "对比 · 测试", "个", os.path.join(tmp, "c.png")),
                charts.combo(["1月", "2月"], [100, 120], [0.95, 1.02], "组合 · 测试", "万元", "达成率", os.path.join(tmp, "d.png"), rate=True),
            ]
            for o in outs:
                self.assertTrue(os.path.getsize(o) > 5000, o)
        finally:
            shutil.rmtree(tmp)

    def test_tables(self):
        import pandas as pd
        import openpyxl
        from tables import save_excel
        tmp = tempfile.mkdtemp()
        try:
            p = save_excel(os.path.join(tmp, "t.xlsx"),
                           {"汇总": pd.DataFrame({"仓库": ["A仓", "B仓"], "金额": [1234.5, 99], "占比": [0.9, 0.1]})},
                           titles={"汇总": "测试汇总"}, notes={"汇总": ["结论 1"]})
            wb = openpyxl.load_workbook(p)
            ws = wb["汇总"]
            cells = " ".join(str(c.value) for row in ws.iter_rows(max_row=3) for c in row if c.value)
            self.assertIn(MARK, cells)
            self.assertTrue(ws.freeze_panes)
            self.assertGreater(ws.column_dimensions["A"].width or 0, 6)
        finally:
            shutil.rmtree(tmp)


class T05_Stamp(unittest.TestCase):
    def test_stamp_idempotent(self):
        import pandas as pd
        import openpyxl
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, "x.xlsx")
            pd.DataFrame({"a": [1]}).to_excel(p, index=False)
            for _ in range(2):
                r = run(os.path.join(COMPANY, "stamp_internal.py"), tmp, "-r")
                self.assertEqual(r.returncode, 0, r.stderr)
            ws = openpyxl.load_workbook(p).active
            self.assertEqual(ws.oddFooter.right.text.count(MARK), 1)
        finally:
            shutil.rmtree(tmp)


class T06_Office(unittest.TestCase):
    @unittest.skipUnless(shutil.which("soffice"), "没有 LibreOffice")
    def test_pptx_to_pdf(self):
        from pptx import Presentation
        tmp = tempfile.mkdtemp()
        try:
            p = os.path.join(tmp, "测试.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = "中文标题测试"
            prs.save(p)
            r = run(os.path.join(COMPANY, "office2pdf.py"), p)
            self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
            self.assertTrue(os.path.isfile(os.path.join(tmp, "测试.pdf")))
            if shutil.which("pdftotext"):
                txt = subprocess.run(["pdftotext", os.path.join(tmp, "测试.pdf"), "-"], capture_output=True, text=True).stdout
                self.assertIn("中文标题测试", txt)
        finally:
            shutil.rmtree(tmp)

    def test_weekly_help(self):
        self.assertEqual(run(os.path.join(KIT, "weekly.py"), "--help").returncode, 0)


GUARD_SERVER = r"""
import subprocess, sys, time
long = subprocess.Popen("sleep 300 & sleep 301", shell=True, start_new_session=True)
shell = subprocess.Popen(["/bin/bash"], stdin=subprocess.PIPE, start_new_session=True)
short = subprocess.Popen("sleep 1", shell=True, start_new_session=True)
time.sleep(7)
print("long", long.poll(), "shell", shell.poll(), "short", short.poll(), flush=True)
shell.kill()
"""


class T09_TermGuard(unittest.TestCase):
    def test_kill_long_keep_shell_and_clean_tmp(self):
        tmp = tempfile.mkdtemp()
        try:
            old = os.path.join(tmp, "旧临时目录")
            os.makedirs(old)
            new = os.path.join(tmp, "新临时目录")
            os.makedirs(new)
            t = time.time() - 3 * 86400
            os.utime(old, (t, t))
            srv = os.path.join(tmp, "server.py")
            open(srv, "w", encoding="utf-8").write(GUARD_SERVER)
            guard = os.path.join(COMPANY, "term_guard.py")
            sh = f"TERM_CMD_MAX_MIN=0.05 TERM_GUARD_SCAN_SEC=1 TERM_GUARD_TMP={tmp} {PY} {guard} & exec {PY} {srv}"
            r = subprocess.run(["bash", "-c", sh], capture_output=True, text=True, timeout=60)
            self.assertIn("long -15", r.stdout, r.stdout + r.stderr)       # 超时命令被终止
            self.assertIn("shell None", r.stdout)                           # 交互式 shell 不动
            self.assertIn("short 0", r.stdout)                              # 正常短命令不受影响
            self.assertIn("终止超时命令", r.stdout)
            self.assertFalse(os.path.exists(old))                           # 超过 24 小时的临时目录被清理
            self.assertTrue(os.path.exists(new))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    print("公司工具包回归测试（只用 /tmp 临时文件）\n")
    res = unittest.main(verbosity=2, exit=False, warnings="ignore").result
    bad = len(res.failures) + len(res.errors)
    print(f"\n结果：共 {res.testsRun} 项，失败 {bad} 项，跳过 {len(res.skipped)} 项" + ("　全部通过" if not bad else "　有问题，先别同步 NAS"))
    sys.exit(1 if bad else 0)
