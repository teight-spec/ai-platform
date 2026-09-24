# -*- coding: utf-8 -*-
"""
tables.py —— 公司统一 Excel 输出

用法：
    import sys; sys.path.insert(0, "/opt/company/kit")
    from tables import save_excel
    save_excel("汇总.xlsx", {"结论与汇总": df1, "清洗后明细": df2},
               titles={"结论与汇总": "入库金额汇总 · 2026-W38"},
               notes={"结论与汇总": ["结论1……", "结论2……"]})

规则（已内置）：
  * 第 1 行标题（含「内部文件，禁止外传」），可选说明行，然后是表头
  * 表头加粗、蓝底白字、冻结表头、自动筛选
  * 列宽按内容撑开（中文按 2 个字符宽），行高留余量，不会挤在一起
  * 数字列千分位；列名含「率」「占比」「%」的按百分比显示
"""
import datetime as _dt

MARK = "内部文件，禁止外传"
BLUE = "2F6FB0"


def _width(v):
    s = "" if v is None else str(v)
    return sum(2 if ord(ch) > 127 else 1 for ch in s)


def save_excel(path, sheets, titles=None, notes=None, pct_cols=None, decimals=2):
    """sheets: {sheet名: pandas.DataFrame}（按顺序写入）。返回 path。"""
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    titles, notes, pct_cols = titles or {}, notes or {}, set(pct_cols or [])
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for name, df in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        df = df.copy() if isinstance(df, pd.DataFrame) else pd.DataFrame(df)
        ncol = max(len(df.columns), 1)
        title = titles.get(name, str(name))
        ws.cell(1, 1, f"{title}　【{MARK}】").font = Font(size=14, bold=True, color="262626")
        ws.row_dimensions[1].height = 28
        r = 2
        for line in notes.get(name, []) or []:
            c = ws.cell(r, 1, line)
            c.font = Font(size=10.5, color="404040")
            c.alignment = Alignment(wrap_text=False, vertical="center")
            ws.row_dimensions[r].height = 20
            r += 1
        r += 0 if r == 2 else 1
        header_row = r
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(header_row, j, str(col))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=BLUE)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        ws.row_dimensions[header_row].height = 26
        is_pct = [(str(col) in pct_cols) or any(k in str(col) for k in ("率", "占比", "%")) for col in df.columns]
        for i, row in enumerate(df.itertuples(index=False), header_row + 1):
            for j, v in enumerate(row, 1):
                if v is not None and not isinstance(v, str):
                    try:
                        if pd.isna(v):
                            v = None
                    except (TypeError, ValueError):
                        pass
                if isinstance(v, pd.Timestamp):
                    v = v.to_pydatetime()
                c = ws.cell(i, j, v)
                c.border = border
                c.alignment = Alignment(vertical="center")
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    c.number_format = "0.0%" if is_pct[j - 1] else (
                        "#,##0" if float(v).is_integer() else f"#,##0.{'0' * decimals}")
                elif isinstance(v, (_dt.datetime, _dt.date)):
                    c.number_format = "yyyy-mm-dd"
            ws.row_dimensions[i].height = 20
        for j, col in enumerate(df.columns, 1):
            vals = [col] + df.iloc[:200, j - 1].tolist()
            w = max(_width(v if not isinstance(v, float) else f"{v:,.2f}") for v in vals)
            ws.column_dimensions[get_column_letter(j)].width = min(max(w + 4, 10), 60)
        if len(df.columns):
            ws.freeze_panes = ws.cell(header_row + 1, 1)
            ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncol)}{header_row + max(len(df), 1)}"
        ws.oddFooter.right.text = MARK
        ws.sheet_view.showGridLines = False
    wb.save(path)
    return path
