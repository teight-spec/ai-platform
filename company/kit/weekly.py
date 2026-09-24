#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
weekly.py —— 数字表格型周报的固定出数工具（可选）

原则：数字、图表、PPT 版式全部由本脚本按配置固定生成；AI 只负责写「结论」这几句话，
      并且结论里出现的数字必须能在本周数据里找到（render 时自动校验），保证准确、可追溯。

一次性准备（每个部门做一次）：
    python3 /opt/company/kit/weekly.py draft  --dept-dir ~/资材 --file 01_原始数据/出入库.xlsx
        → 根据样本文件生成配置草稿 05_模板/周报配置.json（和用户逐项确认后再用）
    python3 /opt/company/kit/weekly.py check  --dept-dir ~/资材
        → 检查配置：文件找不找得到、列名对不对、每周有多少行

每周三步：
    ① python3 /opt/company/kit/weekly.py prepare --dept-dir ~/资材 [--week 2026-W38]
        → 03_周报/2026-W38/：数据.xlsx、图表、facts.md（给 AI 看的数字摘要）、结论.json（待填）
    ② AI 读 facts.md，只改写 结论.json（核心结论 1~3 条、异常说明、下周关注）
    ③ python3 /opt/company/kit/weekly.py render  --dept-dir ~/资材 --week 2026-W38
        → 校验结论里的数字 → 生成 PPT → 转 PDF（PPT 移到「底稿」）→ 加内部标注
        → 交付：03_周报/2026-W38/周报_<部门>_2026-W38.pdf

不写 --week 时：周五、周六、周日默认本周，周一到周四默认上周。
"""
import argparse
import ast
import datetime as dt
import glob
import json
import math
import os
import re
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

CONFIG_REL = os.path.join("05_模板", "周报配置.json")
HISTORY_REL = os.path.join("03_周报", "指标历史.csv")
MARK = "内部文件，禁止外传"
AGGS = {"求和", "计数", "去重计数", "平均", "最大", "最小"}
BLUE, GREEN, ORANGE, GRAY, DARK, LIGHT = "2F6FB0", "3A9A5B", "D9822B", "8C8C8C", "262626", "F5F7FA"


class ConfigError(Exception):
    pass


# ============================================================ 周次
def week_of(d):
    y, w, _ = d.isocalendar()
    return f"{y}-W{w:02d}"


def week_range(week):
    m = re.fullmatch(r"(\d{4})-?W(\d{1,2})", week.strip().upper())
    if not m:
        raise ConfigError(f"周次格式不对：{week}（应为 2026-W38）")
    mon = dt.date.fromisocalendar(int(m.group(1)), int(m.group(2)), 1)
    return mon, mon + dt.timedelta(days=6)


def norm_week(week):
    mon, _ = week_range(week)
    return week_of(mon)


def default_week(today=None):
    today = today or dt.date.today()
    d = today if today.weekday() >= 4 else today - dt.timedelta(days=7)
    return week_of(d)


def weeks_back(week, n):
    mon, _ = week_range(week)
    return [week_of(mon - dt.timedelta(weeks=i)) for i in range(n - 1, -1, -1)]


# ============================================================ 配置
def load_config(dept_dir, path=None):
    path = path or os.path.join(dept_dir, CONFIG_REL)
    if not os.path.isfile(path):
        raise ConfigError(f"找不到周报配置：{path}\n先运行 draft 生成配置草稿，和用户确认后保存到这个位置。")
    try:
        with open(path, encoding="utf-8") as f:
            cfg = json.load(f)
    except json.JSONDecodeError as e:
        raise ConfigError(f"周报配置不是合法 JSON（第 {e.lineno} 行第 {e.colno} 列）：{e.msg}")
    errs = []
    for k in ("标题", "数据源", "指标"):
        if k not in cfg:
            errs.append(f"缺少「{k}」")
    srcs = cfg.get("数据源") or {}
    names = set()
    for m in cfg.get("指标") or []:
        n = m.get("名称")
        if not n:
            errs.append(f"有指标没写「名称」：{m}")
            continue
        if n in names:
            errs.append(f"指标名称重复：{n}")
        names.add(n)
        if "公式" in m:
            continue
        if m.get("数据源") not in srcs:
            errs.append(f"指标「{n}」的数据源「{m.get('数据源')}」不在「数据源」里")
        if m.get("计算") not in AGGS:
            errs.append(f"指标「{n}」的计算方式「{m.get('计算')}」不支持（可用：{'、'.join(sorted(AGGS))}）")
        if m.get("计算") != "计数" and not m.get("列"):
            errs.append(f"指标「{n}」要写「列」（计算方式为计数时可不写）")
    for m in cfg.get("指标") or []:
        if "公式" in m:
            try:
                refs = formula_names(m["公式"])
            except ConfigError as e:
                errs.append(f"指标「{m.get('名称')}」：{e}")
                continue
            miss = [r for r in refs if r not in names]
            if miss:
                errs.append(f"指标「{m.get('名称')}」的公式引用了不存在的指标：{'、'.join(miss)}")
    for g in cfg.get("分组对比") or []:
        base = next((m for m in cfg.get("指标") or [] if m.get("名称") == g.get("指标")), None)
        if not base:
            errs.append(f"分组对比「{g.get('名称')}」的指标「{g.get('指标')}」不存在")
        elif "公式" in base:
            errs.append(f"分组对比「{g.get('名称')}」不能用公式指标，请用求和/计数类指标")
        if not g.get("分组列"):
            errs.append(f"分组对比「{g.get('名称')}」缺少「分组列」")
    if errs:
        raise ConfigError("周报配置有问题：\n  - " + "\n  - ".join(errs))
    cfg.setdefault("趋势周数", 8)
    cfg.setdefault("异常阈值", 0.2)
    cfg.setdefault("部门", os.path.basename(os.path.normpath(dept_dir)))
    return cfg


def formula_names(expr):
    try:
        tree = ast.parse(expr, mode="eval")
    except SyntaxError:
        raise ConfigError(f"公式写法不对：{expr}")
    out = []
    for node in ast.walk(tree):
        if isinstance(node, ast.Name):
            out.append(node.id)
        elif not isinstance(node, (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Load,
                                   ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd)):
            raise ConfigError(f"公式只能用 + - * / 和括号：{expr}")
    return out


def eval_formula(expr, values):
    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant):
            return float(node.value)
        if isinstance(node, ast.Name):
            v = values.get(node.id)
            if v is None:
                raise ZeroDivisionError
            return float(v)
        if isinstance(node, ast.UnaryOp):
            v = ev(node.operand)
            return -v if isinstance(node.op, ast.USub) else v
        if isinstance(node, ast.BinOp):
            a, b = ev(node.left), ev(node.right)
            if isinstance(node.op, ast.Add):
                return a + b
            if isinstance(node.op, ast.Sub):
                return a - b
            if isinstance(node.op, ast.Mult):
                return a * b
            if b == 0:
                raise ZeroDivisionError
            return a / b
        raise ZeroDivisionError
    try:
        return ev(ast.parse(expr, mode="eval"))
    except ZeroDivisionError:
        return None


# ============================================================ 读数据
def find_files(dept_dir, pattern):
    pat = pattern if os.path.isabs(pattern) else os.path.join(dept_dir, pattern)
    files = sorted(set(glob.glob(pat, recursive=True)))
    return [f for f in files if os.path.isfile(f) and not os.path.basename(f).startswith(("~$", "."))
            and "03_周报" not in os.path.relpath(f, dept_dir).split(os.sep)[:1]]


def to_num(s):
    import pandas as pd
    if s.dtype.kind in "if":
        return s.astype(float), 0
    t = s.astype(str).str.replace(",", "", regex=False).str.replace("，", "", regex=False).str.strip()
    t = t.where(~t.isin(["", "nan", "None", "NaT", "-", "—"]))
    num = pd.to_numeric(t, errors="coerce")
    bad = int((num.isna() & t.notna()).sum())
    return num, bad


def load_sources(cfg, dept_dir):
    import pandas as pd
    out, meta, issues = {}, [], []
    for name, src in cfg["数据源"].items():
        pattern = src.get("文件")
        if not pattern:
            raise ConfigError(f"数据源「{name}」缺少「文件」")
        files = find_files(dept_dir, pattern)
        if not files:
            raise ConfigError(f"数据源「{name}」找不到文件：{pattern}（相对于 {dept_dir}）")
        hdr = int(src.get("表头行", 1)) - 1
        sheet = src.get("sheet", 0)
        date_col = src.get("日期列")
        if not date_col:
            raise ConfigError(f"数据源「{name}」缺少「日期列」（周报要按日期分周）")
        frames = []
        for f in files:
            ext = os.path.splitext(f)[1].lower()
            try:
                if ext in (".csv", ".txt"):
                    df = None
                    for enc in ("utf-8-sig", "gbk", "gb18030"):
                        try:
                            df = pd.read_csv(f, header=hdr, encoding=enc, dtype=object)
                            break
                        except UnicodeDecodeError:
                            continue
                    if df is None:
                        raise ConfigError(f"{f} 编码无法识别")
                else:
                    df = pd.read_excel(f, sheet_name=sheet, header=hdr, dtype=object,
                                       engine="xlrd" if ext == ".xls" else None)
            except ValueError as e:
                if "Worksheet" in str(e) or "sheet" in str(e).lower():
                    raise ConfigError(f"数据源「{name}」：文件 {os.path.basename(f)} 里没有 sheet「{sheet}」")
                raise
            df.columns = [str(c).strip() for c in df.columns]
            df = df.dropna(how="all")
            if date_col not in df.columns:
                raise ConfigError(f"数据源「{name}」：文件 {os.path.basename(f)} 没有日期列「{date_col}」。"
                                  f"现有列：{'、'.join(list(df.columns)[:20])}")
            df["_文件"] = os.path.relpath(f, dept_dir)
            frames.append(df)
        df = pd.concat(frames, ignore_index=True)
        dates = pd.to_datetime(df[date_col], errors="coerce")
        bad_dates = int(dates.isna().sum())
        if bad_dates:
            issues.append(f"数据源「{name}」有 {bad_dates} 行日期无法识别，已排除（列「{date_col}」）")
        df = df[dates.notna()].copy()
        df["_日期"] = dates[dates.notna()].dt.date
        df["_周"] = [week_of(d) for d in df["_日期"]]
        cols = [c for c in df.columns if not c.startswith("_")]
        dup = int(df.duplicated(subset=cols).sum())
        if dup and len(files) > 1:
            issues.append(f"数据源「{name}」在不同文件之间有 {dup} 行完全相同，可能重复放了文件（未自动去重，请核对）")
        elif dup:
            issues.append(f"数据源「{name}」有 {dup} 行完全相同（未自动去重，请核对是否为重复录入）")
        out[name] = df
        for f in files:
            sub = df[df["_文件"] == os.path.relpath(f, dept_dir)]
            meta.append({"数据源": name, "文件": os.path.relpath(f, dept_dir), "有效行数": int(len(sub)),
                         "日期范围": f"{sub['_日期'].min()} ~ {sub['_日期'].max()}" if len(sub) else "无有效日期",
                         "修改时间": dt.datetime.fromtimestamp(os.path.getmtime(f)).strftime("%Y-%m-%d %H:%M")})
    return out, meta, issues


def apply_filter(df, flt, label):
    if not flt:
        return df
    mask = None
    for col, cond in flt.items():
        if col not in df.columns:
            raise ConfigError(f"{label} 的筛选列「{col}」不存在")
        s = df[col].astype(str).str.strip()
        if isinstance(cond, dict):
            if "包含" in cond:
                m = s.str.contains(str(cond["包含"]), regex=False)
            elif "不等于" in cond:
                v = cond["不等于"]
                m = ~s.isin([str(x) for x in (v if isinstance(v, list) else [v])])
            else:
                raise ConfigError(f"{label} 的筛选写法不支持：{cond}（可用：值、[值…]、{{\"包含\": …}}、{{\"不等于\": …}}）")
        else:
            m = s.isin([str(x).strip() for x in (cond if isinstance(cond, list) else [cond])])
        mask = m if mask is None else (mask & m)
    return df[mask]


def aggregate(df, m, label):
    how = m["计算"]
    if how == "计数":
        return float(len(df)), 0
    col = m["列"]
    if col not in df.columns:
        raise ConfigError(f"{label}：列「{col}」不存在。现有列：{'、'.join(c for c in df.columns if not c.startswith('_'))}")
    if how == "去重计数":
        s = df[col].dropna().astype(str).str.strip()
        return float(s[s != ""].nunique()), 0
    num, bad = to_num(df[col])
    if len(num.dropna()) == 0:
        return (0.0 if how == "求和" else None), bad
    v = {"求和": num.sum, "平均": num.mean, "最大": num.max, "最小": num.min}[how]()
    return float(v), bad


def scale(m, v):
    if v is None:
        return None
    k = m.get("换算") or 1
    v = v / k
    if m.get("显示") == "百分比":
        return round(v, 6)
    return round(v, int(m.get("小数", 2)))


# ============================================================ 计算
def read_history(dept_dir):
    import csv
    p = os.path.join(dept_dir, HISTORY_REL)
    h = {}
    if os.path.isfile(p):
        with open(p, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                try:
                    h[(r["周次"], r["指标"])] = float(r["值"])
                except (KeyError, ValueError):
                    pass
    return h


def write_history(dept_dir, week, metrics):
    import csv
    p = os.path.join(dept_dir, HISTORY_REL)
    rows = []
    if os.path.isfile(p):
        with open(p, encoding="utf-8-sig") as f:
            rows = [r for r in csv.DictReader(f) if r.get("周次") != week]
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    for m in metrics:
        if m["本周"] is not None:
            rows.append({"周次": week, "指标": m["名称"], "值": m["本周"], "单位": m["单位"], "生成时间": now})
    rows.sort(key=lambda r: (r["周次"], r["指标"]))
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["周次", "指标", "值", "单位", "生成时间"])
        w.writeheader()
        w.writerows(rows)


def compute(cfg, dept_dir, week):
    sources, meta, issues = load_sources(cfg, dept_dir)
    n = max(2, int(cfg["趋势周数"]))
    weeks = weeks_back(week, n)
    prev = weeks[-2]
    mon, sun = week_range(week)
    history = read_history(dept_dir)
    coverage = {name: set(df["_周"]) for name, df in sources.items()}
    for name, df in sources.items():
        wk = df[df["_周"] == week]
        if len(wk) == 0:
            issues.append(f"数据源「{name}」本周（{week}）没有任何数据行，请确认文件是否已放入")
        else:
            days_now = wk["_日期"].nunique()
            before = [df[df["_周"] == w]["_日期"].nunique() for w in weeks_back(week, 5)[:-1]]
            before = sorted(x for x in before if x)
            typical = before[len(before) // 2] if before else 0
            if typical and days_now < typical:
                issues.append(f"数据源「{name}」本周只有 {days_now} 天有数据（往常每周 {typical} 天，最后一天 {max(wk['_日期'])}），"
                              f"本周可能不完整")

    if all((df["_周"] == week).sum() == 0 for df in sources.values()):
        last = max(max(df["_周"]) for df in sources.values() if len(df))
        raise ConfigError(f"{week}（{mon} ~ {sun}）在所有数据源里都没有数据。请先把本周的文件放进 01_原始数据；"
                          f"现有数据最新到 {last}。")
    per_week = {}   # 指标名 -> {周: 值(已换算、按小数位取整)}
    raw_week = {}   # 指标名 -> {周: 未换算的原值}（公式指标用它算，避免取整误差）
    used_hist = []
    bad_nums = {}
    for m in cfg["指标"]:
        if "公式" in m:
            continue
        name, src = m["名称"], m["数据源"]
        df = apply_filter(sources[src], m.get("筛选"), f"指标「{name}」")
        vals, raws = {}, {}
        for w in weeks:
            if w in coverage[src]:
                v, bad = aggregate(df[df["_周"] == w], m, f"指标「{name}」")
                if bad and w == week:
                    bad_nums[name] = bad
                vals[w], raws[w] = scale(m, v), v
            elif (w, name) in history:
                vals[w] = history[(w, name)]
                raws[w] = vals[w] * (m.get("换算") or 1)
                used_hist.append(f"{name}·{w}")
            else:
                vals[w] = raws[w] = None
        per_week[name], raw_week[name] = vals, raws
    for m in cfg["指标"]:
        if "公式" not in m:
            continue
        refs = formula_names(m["公式"])
        vals, raws = {}, {}
        for w in weeks:
            v = eval_formula(m["公式"], {r: raw_week.get(r, {}).get(w) for r in refs})
            vals[w], raws[w] = scale(m, v), v
        per_week[m["名称"]], raw_week[m["名称"]] = vals, raws
    for name, bad in bad_nums.items():
        issues.append(f"指标「{name}」本周有 {bad} 个值不是数字，已按空值处理（请核对原始数据）")
    if used_hist:
        issues.append("以下数值来自以往周报的历史记录（原始数据里已没有这些周）：" + "、".join(used_hist[:10])
                      + ("等" if len(used_hist) > 10 else ""))

    thr = float(cfg["异常阈值"])
    thr_pp = float(cfg.get("百分点阈值", 5)) / 100
    metrics, anomalies = [], []
    for m in cfg["指标"]:
        name = m["名称"]
        cur, last = per_week[name][week], per_week[name][prev]
        chg = None if cur is None or last is None else cur - last
        rate = None if chg is None or not last else chg / abs(last)
        direction = m.get("方向", "越大越好")
        target = m.get("目标")
        judge = "持平"
        moved = (chg is not None and abs(chg) >= 0.0005) if m.get("显示") == "百分比" else (rate is not None and abs(rate) >= 0.005)
        if moved:
            up = chg > 0
            judge = ("好转" if up else "变差") if direction == "越大越好" else \
                    ("变差" if up else "好转") if direction == "越小越好" else ("上升" if up else "下降")
        hit = None
        if target is not None and cur is not None:
            hit = cur >= target if direction != "越小越好" else cur <= target
        item = {"名称": name, "单位": m.get("单位", ""), "显示": m.get("显示", "数值"), "小数": int(m.get("小数", 2)),
                "方向": direction, "本周": cur, "上周": last, "变化": None if chg is None else round(chg, 6),
                "变化率": None if rate is None else round(rate, 6), "目标": target, "达标": hit, "判定": judge,
                "口径": describe(m), "出图": m.get("出图", True), "趋势": {"周": weeks, "值": [per_week[name][w] for w in weeks]}}
        metrics.append(item)
        if item["显示"] == "百分比":
            if chg is not None and abs(chg) >= thr_pp:
                anomalies.append({"指标": name, "类型": "波动",
                                  "说明": f"较上周 {chg * 100:+.1f}个百分点（阈值 ±{thr_pp * 100:.0f}个百分点）"})
        elif rate is not None and abs(rate) >= thr:
            anomalies.append({"指标": name, "类型": "波动", "说明": f"较上周 {rate * 100:+.1f}%（阈值 ±{thr * 100:.0f}%）"})
        if hit is False:
            anomalies.append({"指标": name, "类型": "未达标",
                              "说明": f"本周 {show(item, cur)}，目标 {show(item, target)}"})
        if cur is None:
            anomalies.append({"指标": name, "类型": "无数据", "说明": "本周没有数据"})

    groups = []
    for g in cfg.get("分组对比") or []:
        m = next(x for x in cfg["指标"] if x["名称"] == g["指标"])
        df = apply_filter(sources[m["数据源"]], m.get("筛选"), f"分组对比「{g.get('名称')}」")
        col = g["分组列"]
        if col not in df.columns:
            raise ConfigError(f"分组对比「{g.get('名称')}」的分组列「{col}」不存在")
        rows = []
        keys = df[df["_周"].isin([week, prev])][col].fillna("（空）").astype(str).str.strip().unique()
        for k in keys:
            sub = df[df[col].fillna("（空）").astype(str).str.strip() == k]
            c = scale(m, aggregate(sub[sub["_周"] == week], m, "分组")[0])
            p = scale(m, aggregate(sub[sub["_周"] == prev], m, "分组")[0])
            d = None if c is None or p is None else round(c - p, 6)
            rows.append({"项": k, "本周": c, "上周": p, "变化": d,
                         "变化率": None if d is None or not p else round(d / abs(p), 6)})
        rows.sort(key=lambda r: -(r["本周"] or 0))
        top = int(g.get("前N", 8))
        shown, rest = rows[:top], rows[top:]
        other = None
        if rest:
            other = {"项": f"其他（{len(rest)} 项）", "本周": round(sum(r["本周"] or 0 for r in rest), 6),
                     "上周": round(sum(r["上周"] or 0 for r in rest), 6)}
            other["变化"] = round(other["本周"] - other["上周"], 6)
            other["变化率"] = None if not other["上周"] else round(other["变化"] / abs(other["上周"]), 6)
        movers = sorted([r for r in rows if r["变化"] is not None], key=lambda r: -abs(r["变化"]))[:3]
        groups.append({"名称": g.get("名称") or f"按{col}的{m['名称']}", "分组列": col, "指标": m["名称"],
                       "单位": m.get("单位", ""), "小数": int(m.get("小数", 2)), "显示": m.get("显示", "数值"),
                       "明细": shown, "其他": other, "变化最大": movers})

    return {
        "标题": cfg["标题"], "部门": cfg["部门"], "汇报对象": cfg.get("汇报对象", ""),
        "周次": week, "上周": prev, "日期范围": f"{mon} ~ {sun}",
        "生成时间": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "指标": metrics, "分组": groups, "异常": anomalies, "数据质量": issues, "数据来源": meta,
        "异常阈值": thr, "关注指标": cfg.get("关注指标") or [m["名称"] for m in metrics][:6],
    }


def describe(m):
    if "公式" in m:
        s = f"公式：{m['公式']}"
    else:
        s = f"{m['数据源']}·{m['计算']}" + (f"「{m['列']}」" if m.get("列") else "")
    if m.get("筛选"):
        s += "，筛选 " + "；".join(f"{k}={v}" for k, v in m["筛选"].items())
    if (m.get("换算") or 1) != 1:
        s += f"，÷{m['换算']}"
    return s


def show(item, v, signed=False):
    """按指标的单位/小数显示数值"""
    if v is None:
        return "—"
    if item.get("显示") == "百分比":
        return f"{v * 100:+.1f}%" if signed else f"{v * 100:.1f}%"
    d = int(item.get("小数", 2))
    s = f"{v:+,.{d}f}" if signed else f"{v:,.{d}f}"
    return s + (item.get("单位") or "")


def show_rate(r):
    return "—" if r is None else f"{r * 100:+.1f}%"


def show_change(item):
    """较上周变化：百分比类指标用「个百分点」，其余用变化率"""
    if item.get("显示") == "百分比":
        return "—" if item.get("变化") is None else f"{item['变化'] * 100:+.1f}个百分点"
    return show_rate(item.get("变化率"))


# ============================================================ prepare 输出
def make_charts(facts, out_dir):
    from charts import bar_compare, line
    img = os.path.join(out_dir, "图")
    os.makedirs(img, exist_ok=True)
    paths = {"趋势": {}, "分组": {}}
    for i, m in enumerate(facts["指标"], 1):
        vals = m["趋势"]["值"]
        if all(v is None for v in vals) or m.get("出图") is False:
            continue
        weeks = [w.split("-")[1] for w in m["趋势"]["周"]]
        pctmode = m["显示"] == "百分比"
        series = [None if v is None else (v * 100 if pctmode else v) for v in vals]
        unit = "%" if pctmode else m["单位"]
        p = os.path.join(img, f"趋势_{i:02d}.png")
        line(weeks, {m["名称"]: series}, f"{m['名称']} · 近{len(weeks)}周", unit, p,
             decimals=1 if pctmode else m["小数"], size=(6.4, 4.6),
             target=None if m["目标"] is None else (m["目标"] * 100 if pctmode else m["目标"]))
        paths["趋势"][m["名称"]] = os.path.relpath(p, out_dir)
    for i, g in enumerate(facts["分组"], 1):
        rows = g["明细"] + ([g["其他"]] if g["其他"] else [])
        if not rows:
            continue
        pctmode = g["显示"] == "百分比"
        k = 100 if pctmode else 1
        p = os.path.join(img, f"分组_{i:02d}.png")
        bar_compare([r["项"] for r in rows],
                    {f"上周 {facts['上周'].split('-')[1]}": [(r["上周"] or 0) * k for r in rows],
                     f"本周 {facts['周次'].split('-')[1]}": [(r["本周"] or 0) * k for r in rows]},
                    f"{g['名称']} · 本周 vs 上周", "%" if pctmode else g["单位"], p,
                    decimals=1 if pctmode else g["小数"], size=(8.6, 4.4))
        paths["分组"][g["名称"]] = os.path.relpath(p, out_dir)
    facts["图表"] = paths


def write_workbook(facts, out_dir):
    import pandas as pd
    from tables import save_excel
    ms = facts["指标"]
    summary = pd.DataFrame([{"指标": m["名称"], "单位": "%" if m["显示"] == "百分比" else m["单位"],
                             "本周": None if m["本周"] is None else (m["本周"] * 100 if m["显示"] == "百分比" else m["本周"]),
                             "上周": None if m["上周"] is None else (m["上周"] * 100 if m["显示"] == "百分比" else m["上周"]),
                             "变化率": None if m["显示"] == "百分比" else m["变化率"],
                             "变化(个百分点)": (None if m["变化"] is None else round(m["变化"] * 100, 2)) if m["显示"] == "百分比" else None,
                             "判定": m["判定"],
                             "目标": None if m["目标"] is None else (m["目标"] * 100 if m["显示"] == "百分比" else m["目标"]),
                             "是否达标": {True: "达标", False: "未达标", None: ""}[m["达标"]],
                             "口径": m["口径"]} for m in ms])
    weeks = ms[0]["趋势"]["周"] if ms else []
    trend = pd.DataFrame([{"指标": m["名称"], **{w: v for w, v in zip(m["趋势"]["周"], m["趋势"]["值"])}} for m in ms],
                         columns=["指标"] + weeks)
    grows = []
    for g in facts["分组"]:
        for r in g["明细"] + ([g["其他"]] if g["其他"] else []):
            grows.append({"分组": g["名称"], "项": r["项"], "本周": r["本周"], "上周": r["上周"], "变化": r["变化"], "变化率": r["变化率"]})
    sheets = {"指标汇总": summary, "近期趋势": trend,
              "分组对比": pd.DataFrame(grows, columns=["分组", "项", "本周", "上周", "变化", "变化率"]),
              "异常": pd.DataFrame(facts["异常"], columns=["指标", "类型", "说明"]),
              "数据质量": pd.DataFrame({"事项": facts["数据质量"] or ["未发现问题"]}),
              "数据来源": pd.DataFrame(facts["数据来源"])}
    path = os.path.join(out_dir, f"数据_{facts['部门']}_{facts['周次']}.xlsx")
    save_excel(path, sheets, titles={"指标汇总": f"{facts['标题']} · {facts['周次']}（{facts['日期范围']}）"},
               notes={"指标汇总": [f"生成时间 {facts['生成时间']}；百分比类指标的本周/上周已乘 100。本表由 weekly.py 按配置自动计算，可追溯到「数据来源」。"]},
               pct_cols=["变化率"])
    return path


def write_facts_md(facts, out_dir):
    L = [f"# {facts['标题']} · {facts['周次']}（{facts['日期范围']}）数字摘要",
         "", "写结论只能用本文件里的数字；数字照抄（含单位），不要自己换算或估算。",
         "百分比类指标的变化用「个百分点」表示。", "", "## 指标（本周 / 上周 / 较上周 / 判定）"]
    for m in facts["指标"]:
        t = f"，目标 {show(m, m['目标'])}（{'达标' if m['达标'] else '未达标'}）" if m["目标"] is not None else ""
        L.append(f"- {m['名称']}：{show(m, m['本周'])} / {show(m, m['上周'])} / {show_change(m)} / {m['判定']}{t}")
    L.append("\n## 近期趋势")
    for m in facts["指标"]:
        L.append(f"- {m['名称']}：" + "，".join(f"{w.split('-')[1]} {show(m, v)}" for w, v in zip(m["趋势"]["周"], m["趋势"]["值"])))
    for g in facts["分组"]:
        L.append(f"\n## {g['名称']}（本周 / 上周 / 变化率）")
        for r in g["明细"] + ([g["其他"]] if g["其他"] else []):
            L.append(f"- {r['项']}：{show(g, r['本周'])} / {show(g, r['上周'])} / {show_rate(r['变化率'])}")
        if g["变化最大"]:
            L.append("- 变化最大：" + "；".join(f"{r['项']} {show(g, r['变化'], True)}" for r in g["变化最大"]))
    L.append("\n## 异常（需要在 结论.json 的「异常说明」里逐条写原因推测）")
    L += [f"- {a['指标']}【{a['类型']}】{a['说明']}" for a in facts["异常"]] or ["- 无"]
    L.append("\n## 数据质量（会原样放进 PPT 的「数据说明」页）")
    L += [f"- {x}" for x in facts["数据质量"]] or ["- 未发现问题"]
    L.append("\n## 数据来源")
    L += [f"- {s['数据源']}：{s['文件']}，{s['有效行数']} 行，{s['日期范围']}" for s in facts["数据来源"]]
    L += ["", "## 下一步", "编辑同目录的 结论.json：", "- 核心结论：1~3 条，每条一句话（不超过 40 字），先说结果再说原因；",
          "- 异常说明：上面每个异常指标写一句原因推测，不确定就写「原因待确认」；",
          "- 下周关注：不超过 3 条。", "改完运行 render。"]
    p = os.path.join(out_dir, "facts.md")
    with open(p, "w", encoding="utf-8") as f:
        f.write("\n".join(L) + "\n")
    return p


def cmd_prepare(a):
    week = norm_week(a.week) if a.week else default_week()
    cfg = load_config(a.dept_dir, a.config)
    facts = compute(cfg, a.dept_dir, week)
    out_dir = os.path.join(a.dept_dir, "03_周报", week)
    os.makedirs(out_dir, exist_ok=True)
    make_charts(facts, out_dir)
    xlsx = write_workbook(facts, out_dir)
    with open(os.path.join(out_dir, "facts.json"), "w", encoding="utf-8") as f:
        json.dump(facts, f, ensure_ascii=False, indent=1, default=str)
    with open(os.path.join(out_dir, "配置快照.json"), "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=1)
    concl = os.path.join(out_dir, "结论.json")
    if not os.path.exists(concl) or a.reset:
        tpl = {"核心结论": ["", "", ""],
               "异常说明": {x["指标"]: "原因待确认" for x in facts["异常"]},
               "下周关注": [""]}
        with open(concl, "w", encoding="utf-8") as f:
            json.dump(tpl, f, ensure_ascii=False, indent=1)
    md = write_facts_md(facts, out_dir)
    write_history(a.dept_dir, week, facts["指标"])
    print(f"周次：{week}（{facts['日期范围']}），对比上周 {facts['上周']}")
    print(f"输出目录：{out_dir}")
    print(f"  数据表：{os.path.basename(xlsx)}；图表 {len(facts['图表']['趋势']) + len(facts['图表']['分组'])} 张")
    print(f"  异常 {len(facts['异常'])} 项；数据质量提示 {len(facts['数据质量'])} 项")
    print(f"下一步：读 {md}，然后改写 {concl}，再运行 render。")


# ============================================================ render：校验
_SKIP = [r"\d{4}\s*-\s*W\d{1,2}", r"\bW\d{1,2}\b", r"\d{4}[-/年.]\d{1,2}[-/月.]\d{1,2}日?", r"\d{1,2}月\d{1,2}日",
         r"\d{1,2}月", r"\d{4}年", r"第\s*\d+\s*[周名位]", r"\d+\s*[月日号]"]


def allowed_numbers(facts):
    vals = []

    def add(v, item=None):
        if v is None:
            return
        v = float(v)
        vals.append(v)
        if item is not None and item.get("显示") == "百分比":
            vals.append(v * 100)

    for m in facts["指标"]:
        for k in ("本周", "上周", "变化", "目标"):
            add(m[k], m)
            if m[k] is not None:
                add(abs(m[k]), m)
        for v in m["趋势"]["值"]:
            add(v, m)
        if m["变化率"] is not None:
            vals += [m["变化率"] * 100, abs(m["变化率"] * 100)]
    for g in facts["分组"]:
        for r in g["明细"] + ([g["其他"]] if g["其他"] else []):
            for k in ("本周", "上周", "变化"):
                add(r[k], g)
                if r[k] is not None:
                    add(abs(r[k]), g)
            if r["变化率"] is not None:
                vals += [r["变化率"] * 100, abs(r["变化率"] * 100)]
    for s in facts["数据来源"]:
        vals.append(float(s["有效行数"]))
    vals.append(facts["异常阈值"] * 100)
    return vals


def check_numbers(texts, facts):
    allowed = allowed_numbers(facts)
    bad = []
    for t in texts:
        s = str(t)
        for p in _SKIP:
            s = re.sub(p, " ", s)
        for mt in re.finditer(r"[-+]?\d[\d,]*(?:\.\d+)?\s*%?", s):
            raw = mt.group(0).strip()
            is_pct = raw.endswith("%")
            try:
                v = abs(float(raw.rstrip("%").strip().replace(",", "")))
            except ValueError:
                continue
            if not is_pct and v <= 12 and float(v).is_integer():
                continue  # 3 条、5 个仓库 这类小整数不校验
            ok = False
            for a in allowed:
                a = abs(a)
                decs = len(raw.rstrip("%").split(".")[1]) if "." in raw else 0
                tol = max(0.5 * 10 ** (-decs), 0.005 * a) + 1e-9
                if abs(v - a) <= tol:
                    ok = True
                    break
            if not ok:
                bad.append((raw, str(t)))
    return bad


def load_conclusions(out_dir, facts):
    p = os.path.join(out_dir, "结论.json")
    if not os.path.isfile(p):
        raise ConfigError(f"找不到 {p}，先运行 prepare")
    try:
        with open(p, encoding="utf-8") as f:
            c = json.load(f)
    except json.JSONDecodeError as e:
        raise ConfigError(f"结论.json 不是合法 JSON（第 {e.lineno} 行）：{e.msg}")
    core = [x.strip() for x in c.get("核心结论") or [] if isinstance(x, str) and x.strip()]
    focus = [x.strip() for x in c.get("下周关注") or [] if isinstance(x, str) and x.strip()]
    notes = {k: str(v).strip() for k, v in (c.get("异常说明") or {}).items()}
    errs = []
    if not 1 <= len(core) <= 3:
        errs.append(f"「核心结论」要写 1~3 条，现在是 {len(core)} 条")
    for x in core:
        if len(x) > 60:
            errs.append(f"核心结论太长（{len(x)} 字，限 60）：{x[:30]}…")
    if len(focus) > 3:
        errs.append("「下周关注」不超过 3 条")
    names = {a["指标"] for a in facts["异常"]}
    for n in names - set(notes):
        notes[n] = "原因待确认"
    bad = check_numbers(core + focus + list(notes.values()), facts)
    for raw, t in bad:
        errs.append(f"数字「{raw}」在本周数据里找不到（出自：{t[:40]}）——请照抄 facts.md 里的数字")
    if errs:
        raise ConfigError("结论.json 需要修改：\n  - " + "\n  - ".join(errs))
    return core, notes, focus


# ============================================================ render：PPT
def _font(run, size=None, bold=None, color=None, name="Microsoft YaHei"):
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    f = run.font
    f.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)
    if size:
        f.size = __import__("pptx.util", fromlist=["Pt"]).Pt(size)
    if bold is not None:
        f.bold = bold
    if color:
        f.color.rgb = RGBColor.from_string(color)


def _text(slide, x, y, w, h, text, size=14, bold=False, color=DARK, align=None, anchor=None, wrap=True):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.05)
    if anchor:
        tf.vertical_anchor = {"middle": MSO_ANCHOR.MIDDLE, "top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "left": PP_ALIGN.LEFT}[align]
        segs = ln if isinstance(ln, list) else [(ln, {})]
        for seg, st in segs:
            r = p.add_run()
            r.text = seg
            _font(r, st.get("size", size), st.get("bold", bold), st.get("color", color))
    return tb


def _rect(slide, x, y, w, h, fill, line=None):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _pic(slide, path, x, y, max_w, max_h):
    from PIL import Image
    from pptx.util import Inches
    with Image.open(path) as im:
        iw, ih = im.size
    r = min(max_w / iw, max_h / ih)
    w, h = iw * r, ih * r
    slide.shapes.add_picture(path, Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2), Inches(w), Inches(h))


def _table(slide, x, y, w, rows, col_w, header_fill=BLUE, size=11, row_h=0.4):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = shp.table
    tot = sum(col_w)
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(w * cw / tot)
    for i in range(nr):
        t.rows[i].height = Inches(row_h)
        for j in range(nc):
            cell = t.cell(i, j)
            val = rows[i][j]
            color = None
            if isinstance(val, tuple):
                val, color = val
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            _font(r, size, i == 0, "FFFFFF" if i == 0 else (color or DARK))
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(header_fill if i == 0 else ("FFFFFF" if i % 2 else LIGHT))
    return shp


class Deck:
    def __init__(self, facts):
        from pptx import Presentation
        from pptx.util import Inches
        self.f = facts
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.n = 0

    def slide(self, title, subtitle=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.n += 1
        _rect(s, 0, 0, 0.14, 7.5, BLUE)
        _text(s, 0.5, 0.3, 10.5, 0.6, title, size=24, bold=True)
        if subtitle:
            _text(s, 0.5, 0.88, 12, 0.4, subtitle, size=12, color="595959")
        _rect(s, 0.5, 1.3, 12.33, 0.03, "D9D9D9")
        foot = f"{self.f['部门']} · {self.f['周次']}（{self.f['日期范围']}）"
        _text(s, 0.5, 7.02, 7, 0.35, foot, size=9, color=GRAY)
        _text(s, 8.3, 7.02, 4.5, 0.35, f"{MARK}　{self.n}", size=9, color=GRAY, align="right")
        return s

    def save(self, path):
        self.prs.save(path)


def render_deck(facts, core, notes, focus, out_dir):
    f = facts
    d = Deck(f)
    ms = {m["名称"]: m for m in f["指标"]}

    # 1 封面
    from pptx.util import Inches  # noqa: F401
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6])
    d.n += 1
    _rect(s, 0, 0, 13.333, 7.5, "FFFFFF")
    _rect(s, 0, 0, 0.35, 7.5, BLUE)
    _rect(s, 1.0, 3.55, 1.2, 0.06, GREEN)
    _text(s, 1.0, 2.2, 11, 1.2, f["标题"], size=40, bold=True)
    _text(s, 1.0, 3.8, 11, 0.6, f"{f['周次']}　{f['日期范围']}", size=20, color="404040")
    sub = f"{f['部门']}" + (f"　汇报对象：{f['汇报对象']}" if f.get("汇报对象") else "") + f"　生成于 {f['生成时间']}"
    _text(s, 1.0, 4.5, 11, 0.5, sub, size=14, color=GRAY)
    _text(s, 1.0, 6.7, 11, 0.4, MARK, size=11, color=ORANGE, bold=True)

    # 2 本周结论 + KPI 卡片
    s = d.slide("本周结论", f"对比上周 {f['上周']}；绿色=好转，橙色=变差")
    y = 1.55
    for i, c in enumerate(core, 1):
        _rect(s, 0.5, y, 0.5, 0.5, BLUE)
        _text(s, 0.5, y, 0.5, 0.5, str(i), size=18, bold=True, color="FFFFFF", align="center", anchor="middle")
        _text(s, 1.15, y - 0.02, 11.6, 0.56, c, size=17, bold=True, anchor="middle")
        y += 0.72
    cards = [ms[n] for n in f["关注指标"] if n in ms][:6]
    if cards:
        n = len(cards)
        gap = 0.2
        w = (12.33 - gap * (n - 1)) / n
        y0 = 4.0
        for i, m in enumerate(cards):
            x = 0.5 + i * (w + gap)
            _rect(s, x, y0, w, 2.7, LIGHT, "E3E7ED")
            _text(s, x + 0.15, y0 + 0.15, w - 0.3, 0.4, m["名称"], size=13, color="595959")
            val = show(m, m["本周"])
            _text(s, x + 0.15, y0 + 0.6, w - 0.3, 0.8, val, size=26 if len(val) <= 9 else 20, bold=True)
            col = GREEN if m["判定"] == "好转" else ORANGE if m["判定"] == "变差" else GRAY
            arrow = "▲" if (m["变化"] or 0) > 0 else "▼" if (m["变化"] or 0) < 0 else "■"
            _text(s, x + 0.15, y0 + 1.45, w - 0.3, 0.4, f"较上周 {arrow} {show_change(m)}", size=13, bold=True, color=col)
            _text(s, x + 0.15, y0 + 1.85, w - 0.3, 0.35, f"上周 {show(m, m['上周'])}", size=11, color=GRAY)
            if m["目标"] is not None:
                _text(s, x + 0.15, y0 + 2.2, w - 0.3, 0.35, f"目标 {show(m, m['目标'])}　{'达标' if m['达标'] else '未达标'}",
                      size=11, color=GREEN if m["达标"] else ORANGE)

    # 3 趋势（每页 2 张）
    trend = [(n, p) for n, p in f["图表"]["趋势"].items()]
    for i in range(0, len(trend), 2):
        pair = trend[i:i + 2]
        s = d.slide("关键指标趋势", f"近 {len(f['指标'][0]['趋势']['周'])} 周，最后一个点为本周")
        for j, (name, p) in enumerate(pair):
            _pic(s, os.path.join(out_dir, p), 0.5 + j * 6.25, 1.55, 6.0 if len(pair) == 2 else 12.3, 5.3)

    # 4 分组对比（每页 1 张 + 变化最大）
    for g in f["分组"]:
        p = f["图表"]["分组"].get(g["名称"])
        if not p:
            continue
        s = d.slide(g["名称"], f"本周 vs 上周；按本周{g['指标']}从高到低")
        _pic(s, os.path.join(out_dir, p), 0.5, 1.55, 8.6, 5.3)
        _rect(s, 9.4, 1.7, 3.43, 4.9, LIGHT, "E3E7ED")
        _text(s, 9.6, 1.85, 3.1, 0.4, "变化最大", size=15, bold=True)
        yy = 2.4
        for r in g["变化最大"]:
            up = (r["变化"] or 0) > 0
            _text(s, 9.6, yy, 3.1, 0.4, r["项"], size=13, bold=True)
            _text(s, 9.6, yy + 0.38, 3.1, 0.4, f"{show(g, r['变化'], True)}（{show_rate(r['变化率'])}）", size=12,
                  color=GREEN if up else ORANGE)
            yy += 0.95
        if not g["变化最大"]:
            _text(s, 9.6, yy, 3.1, 0.4, "无可比数据", size=12, color=GRAY)

    # 5 异常与说明
    s = d.slide("异常与说明", f"波动阈值 ±{f['异常阈值'] * 100:.0f}%；未达标、无数据也列在这里")
    if f["异常"]:
        rows = [["指标", "类型", "本周", "上周", "较上周", "说明"]]
        for a in f["异常"][:10]:
            m = ms.get(a["指标"], {})
            col = ORANGE if a["类型"] != "波动" or m.get("判定") == "变差" else GREEN
            rows.append([a["指标"], (a["类型"], col), show(m, m.get("本周")), show(m, m.get("上周")),
                         (show_change(m) if m else "—", col), notes.get(a["指标"], "原因待确认")])
        _table(s, 0.5, 1.6, 12.33, rows, [2.2, 1.2, 1.7, 1.7, 1.3, 5.2], size=12, row_h=0.46)
        if len(f["异常"]) > 10:
            _text(s, 0.5, 6.6, 12, 0.35, f"另有 {len(f['异常']) - 10} 项见 数据表「异常」页", size=11, color=GRAY)
    else:
        _text(s, 0.5, 3.2, 12.33, 0.8, "本周没有超出阈值的异常", size=24, color=GREEN, bold=True, align="center")

    # 6 下周关注 + 数据说明
    s = d.slide("下周关注与数据说明")
    _text(s, 0.5, 1.55, 6, 0.45, "下周关注", size=17, bold=True, color=BLUE)
    _text(s, 0.5, 2.1, 6, 4.5, [f"{i}. {x}" for i, x in enumerate(focus, 1)] or ["（无）"], size=15)
    _rect(s, 6.9, 1.55, 5.93, 5.2, LIGHT, "E3E7ED")
    _text(s, 7.1, 1.7, 5.5, 0.45, "数据说明（存疑项）", size=15, bold=True, color=ORANGE)
    q = f["数据质量"] or ["未发现数据问题"]
    _text(s, 7.1, 2.2, 5.5, 4.4, [f"· {x}" for x in q[:8]], size=11.5, color="404040")

    # 7 数据来源与口径
    s = d.slide("数据来源与口径", "所有数字由 weekly.py 按「05_模板/周报配置.json」从原始文件计算，可复算")
    rows = [["数据源", "文件", "有效行数", "日期范围"]] + \
           [[x["数据源"], x["文件"], f"{x['有效行数']:,}", x["日期范围"]] for x in f["数据来源"][:6]]
    _table(s, 0.5, 1.55, 12.33, rows, [1.6, 5.5, 1.2, 3.0], size=10.5, row_h=0.36)
    y2 = 1.75 + 0.36 * len(rows)
    rows = [["指标", "口径"]] + [[m["名称"], m["口径"]] for m in f["指标"][:10]]
    _table(s, 0.5, y2, 12.33, rows, [2.2, 10.1], size=10.5, row_h=0.34)

    name = f"周报_{f['部门']}_{f['周次']}.pptx"
    path = os.path.join(out_dir, name)
    d.save(path)
    return path


def cmd_render(a):
    week = norm_week(a.week) if a.week else default_week()
    out_dir = os.path.join(a.dept_dir, "03_周报", week)
    fp = os.path.join(out_dir, "facts.json")
    if not os.path.isfile(fp):
        raise ConfigError(f"{week} 还没有 prepare（找不到 {fp}）")
    with open(fp, encoding="utf-8") as f:
        facts = json.load(f)
    core, notes, focus = load_conclusions(out_dir, facts)
    pptx = render_deck(facts, core, notes, focus, out_dir)
    office = os.path.join(os.path.dirname(HERE), "office2pdf.py")
    stamp = os.path.join(os.path.dirname(HERE), "stamp_internal.py")
    if a.no_pdf:
        print(f"已生成 PPT（未转 PDF）：{pptx}")
        return
    r = subprocess.run([sys.executable, office, pptx], capture_output=True, text=True)
    print(r.stdout.strip())
    if r.returncode != 0:
        raise ConfigError("转 PDF 失败：" + (r.stderr or r.stdout)[-400:])
    pdf = os.path.splitext(pptx)[0] + ".pdf"
    if os.path.isfile(stamp):
        subprocess.run([sys.executable, stamp, out_dir, "-r"], capture_output=True, text=True)
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout \
        if shutil_which("pdftotext") else ""
    pages = txt.count("\f") or "?"
    print(f"完成：{pdf}（{pages} 页）")
    print("交付给用户的只有这个 PDF；数据表 数据_*.xlsx 可作为附件。")


def shutil_which(x):
    import shutil
    return shutil.which(x)


# ============================================================ check / draft
def cmd_check(a):
    cfg = load_config(a.dept_dir, a.config)
    sources, meta, issues = load_sources(cfg, a.dept_dir)
    print("配置格式正确。")
    for s in meta:
        print(f"  [{s['数据源']}] {s['文件']}：{s['有效行数']} 行，{s['日期范围']}")
    for name, df in sources.items():
        wk = df.groupby("_周").size().tail(8)
        print(f"  [{name}] 最近各周行数：" + "，".join(f"{k} {v}" for k, v in wk.items()))
    week = norm_week(a.week) if a.week else default_week()
    facts = compute(cfg, a.dept_dir, week)
    print(f"\n{week} 试算（不写文件）：")
    for m in facts["指标"]:
        print(f"  {m['名称']}：本周 {show(m, m['本周'])}，上周 {show(m, m['上周'])}，{show_change(m)}")
    for x in facts["数据质量"]:
        print(f"  ! {x}")


def cmd_draft(a):
    import pandas as pd
    path = a.file if os.path.isabs(a.file) else os.path.join(a.dept_dir, a.file)
    if not os.path.isfile(path):
        raise ConfigError(f"找不到样本文件：{path}")
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        df = pd.read_csv(path, header=a.header - 1, dtype=object, encoding="utf-8-sig")
        sheet = 0
    else:
        sheet = a.sheet if a.sheet is not None else 0
        if isinstance(sheet, str) and sheet.isdigit():
            sheet = int(sheet)
        df = pd.read_excel(path, sheet_name=sheet, header=a.header - 1, dtype=object,
                           engine="xlrd" if ext == ".xls" else None)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(how="all")
    date_col, nums, cats = None, [], []
    for c in df.columns:
        s = df[c].dropna()
        if not len(s):
            continue
        num, bad = to_num(s)
        if num.notna().mean() >= 0.9:
            nums.append(c)
            continue
        if date_col is None and pd.to_datetime(s, errors="coerce").notna().mean() >= 0.9:
            date_col = c
            continue
        k = s.astype(str).nunique()
        if 2 <= k <= 30:
            cats.append(c)
    if date_col is None:
        raise ConfigError("样本里没找到日期列。周报需要按日期分周，请确认文件里有日期列（或用 --header 指定表头行）")
    rel = os.path.relpath(path, a.dept_dir)
    folder, base = os.path.split(rel)
    stem = re.sub(r"[\d\-_.年月日周Ww]+$", "", os.path.splitext(base)[0]) or os.path.splitext(base)[0]
    pattern = os.path.join(folder, f"*{stem}*{os.path.splitext(base)[1]}")
    src = "数据"
    metrics = [{"名称": "记录数", "数据源": src, "计算": "计数", "单位": "条", "小数": 0, "方向": "无"}]
    for c in nums[:5]:
        metrics.append({"名称": f"{c}合计", "数据源": src, "计算": "求和", "列": c, "单位": "", "小数": 1,
                        "方向": "越大越好", "目标": None})
    groups = [{"名称": f"按{c}对比", "指标": metrics[1]["名称"] if len(metrics) > 1 else "记录数", "分组列": c, "前N": 8}
              for c in cats[:2]]
    cfg = {"标题": f"{os.path.basename(os.path.normpath(a.dept_dir))}周报", "部门": os.path.basename(os.path.normpath(a.dept_dir)),
           "汇报对象": "", "数据源": {src: {"文件": pattern, "sheet": sheet, "表头行": a.header, "日期列": date_col}},
           "指标": metrics, "分组对比": groups, "关注指标": [m["名称"] for m in metrics][:4],
           "趋势周数": 8, "异常阈值": 0.2}
    out = os.path.join(a.dept_dir, CONFIG_REL)
    if os.path.exists(out) and not a.force:
        out = os.path.join(a.dept_dir, "05_模板", "周报配置.草稿.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    print(f"已生成配置草稿：{out}")
    print(f"  日期列：{date_col}；数字列：{'、'.join(nums) or '无'}；可分组的列：{'、'.join(cats) or '无'}")
    print("  请和用户逐项确认：要看哪些指标（求和/计数/去重计数/平均）、单位与换算（如 ÷10000 → 万元）、"
          "方向（越大越好/越小越好/无）、目标值、按什么分组对比、筛选条件。")
    if out.endswith("草稿.json"):
        print(f"  注意：已有正式配置，草稿没有覆盖它；确认后把草稿内容合并进 {CONFIG_REL}")


def main():
    ap = argparse.ArgumentParser(description="数字表格型周报的固定出数工具（可选）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("prepare", "render", "check"):
        p = sub.add_parser(name)
        p.add_argument("--dept-dir", required=True, help="部门文件夹，如 ~/资材")
        p.add_argument("--week", help="周次，如 2026-W38；不写按默认规则")
        p.add_argument("--config", help="配置文件路径（默认 05_模板/周报配置.json）")
        if name == "prepare":
            p.add_argument("--reset", action="store_true", help="重新生成空白 结论.json（会覆盖已写的结论）")
        if name == "render":
            p.add_argument("--no-pdf", action="store_true", help="只生成 PPT，不转 PDF（调试用）")
    p = sub.add_parser("draft")
    p.add_argument("--dept-dir", required=True)
    p.add_argument("--file", required=True, help="样本数据文件（相对部门文件夹）")
    p.add_argument("--sheet")
    p.add_argument("--header", type=int, default=1)
    p.add_argument("--force", action="store_true", help="覆盖已有的正式配置")
    a = ap.parse_args()
    a.dept_dir = os.path.abspath(os.path.expanduser(a.dept_dir))
    try:
        {"prepare": cmd_prepare, "render": cmd_render, "check": cmd_check, "draft": cmd_draft}[a.cmd](a)
    except ConfigError as e:
        print(f"【需要处理】{e}")
        sys.exit(2)


if __name__ == "__main__":
    main()
